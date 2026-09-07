"""Generate real editable Office documents on the GitHub macOS runner."""
from pathlib import Path
import subprocess
import tempfile
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from zipfile import ZipFile
from xml.etree import ElementTree

root = Path("Resources/Templates")
root.mkdir(parents=True, exist_ok=True)
(root / "blank.txt").write_text("", encoding="utf-8")
(root / "blank.md").write_text("", encoding="utf-8")
document = Document()
document.add_paragraph("")
document.save(root / "blank.docx")
workbook = Workbook()
workbook.active.title = "Sheet1"
workbook.save(root / "blank.xlsx")
presentation = Presentation()
presentation.slides.add_slide(presentation.slide_layouts[6])
presentation.save(root / "blank.pptx")
with tempfile.TemporaryDirectory() as tmp:
    rtf = Path(tmp) / "blank.rtf"
    rtf.write_text(r"{\rtf1\ansi\deff0 {\fonttbl {\f0 Helvetica;}}\f0\fs24\par}", encoding="ascii")
    subprocess.run(["textutil", "-convert", "doc", "-output", str(root / "blank.doc"), str(rtf)], check=True)

assert len(Document(root / "blank.docx").paragraphs) >= 1
assert load_workbook(root / "blank.xlsx").sheetnames == ["Sheet1"]
assert len(Presentation(root / "blank.pptx").slides) == 1
for ext, part in [("docx", "word/document.xml"), ("xlsx", "xl/workbook.xml"), ("pptx", "ppt/presentation.xml")]:
    with ZipFile(root / f"blank.{ext}") as archive:
        assert archive.testzip() is None
        assert part in archive.namelist()
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                ElementTree.fromstring(archive.read(name))
assert (root / "blank.doc").read_bytes().startswith(bytes.fromhex("d0cf11e0a1b11ae1"))
subprocess.run(["textutil", "-info", str(root / "blank.doc")], check=True)
print("PASS: all six templates generated; Office containers and XML validated.")
